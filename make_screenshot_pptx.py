"""
Pixel-perfect PPTX: screenshots every HTML slide at 1920x1080,
embeds them as full-bleed images, one per slide.
"""
import asyncio, pathlib, io
from playwright.async_api import async_playwright
from pptx import Presentation
from pptx.util import Emu

BASE   = pathlib.Path(__file__).parent
SLIDES = ["index.html", "slide-1.html", "slide-2.html", "slide-3.html", "slide-4.html"]
OUT    = BASE / "pixel-perfect.pptx"

W_PX, H_PX = 1920, 1080
# Standard widescreen 16:9 in EMU
W_EMU = 9144000
H_EMU = 5143500


async def screenshot_slides():
    images = []
    async with async_playwright() as p:
        browser = await p.chromium.launch()
        page = await browser.new_page(viewport={"width": W_PX, "height": H_PX})
        for html in SLIDES:
            url = f"file://{(BASE / html).resolve()}"
            await page.goto(url, wait_until="networkidle")
            await page.wait_for_timeout(800)   # let web fonts settle
            img_bytes = await page.screenshot(type="png", full_page=False)
            images.append(img_bytes)
            print(f"  captured {html}")
        await browser.close()
    return images


def build_pptx(images):
    prs = Presentation()
    prs.slide_width  = Emu(W_EMU)
    prs.slide_height = Emu(H_EMU)

    blank_layout = prs.slide_layouts[6]   # completely blank

    for i, img_bytes in enumerate(images):
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(
            io.BytesIO(img_bytes),
            left=0, top=0,
            width=Emu(W_EMU), height=Emu(H_EMU),
        )
        print(f"  added slide {i+1}/{len(images)}")

    prs.save(OUT)
    print(f"\nSaved → {OUT}")


async def main():
    print("Taking screenshots…")
    images = await screenshot_slides()
    print("Building PPTX…")
    build_pptx(images)


asyncio.run(main())
